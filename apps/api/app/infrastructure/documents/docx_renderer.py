"""Word (DOCX) layouts of the CV (ATS and Visual) with python-docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.application.ports import DocumentRenderer, RenderedDocument
from app.domain.models import CurriculumVitae
from app.infrastructure.documents.content import DocumentContent
from app.infrastructure.documents.theme import ACCENT_HEX, MUTED_HEX, PRIMARY_HEX, TEXT_HEX

PRIMARY = RGBColor.from_string(PRIMARY_HEX)
ACCENT = RGBColor.from_string(ACCENT_HEX)
TEXT = RGBColor.from_string(TEXT_HEX)
MUTED = RGBColor.from_string(MUTED_HEX)


class DocxCurriculumRenderer(DocumentRenderer):
    def __init__(self, output_path: Path, variant: str, side_margin_pt: int) -> None:
        self._output_path = output_path
        self._variant = variant
        self._side_margin = Pt(side_margin_pt)

    @property
    def variant(self) -> str:
        return self._variant

    @property
    def output_path(self) -> Path:
        return self._output_path

    def render(self, cv: CurriculumVitae, fingerprint: str) -> RenderedDocument:
        content = DocumentContent.from_cv(cv)
        doc = Document()
        for section in doc.sections:
            section.top_margin = section.bottom_margin = Pt(54)
            section.left_margin = section.right_margin = self._side_margin

        props = doc.core_properties
        props.title = f"{content.name} - CV"
        props.author = content.name
        props.subject = "Hoja de vida"
        props.identifier = fingerprint

        self._write(doc, content)
        self._output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(self._output_path))
        return RenderedDocument(self._variant, self._output_path)

    def _write(self, doc: DocxDocument, c: DocumentContent) -> None:
        self._centered(doc, c.name, bold=True, size=18, color=PRIMARY)
        self._centered(doc, c.title, size=10, color=ACCENT)
        for line in c.contact_lines:
            self._centered(doc, line, size=9, color=TEXT)
        self._centered(doc, f"Portafolio: {c.portfolio}", size=9, color=ACCENT)

        self._section(doc, "Perfil Profesional")
        self._run(doc.add_paragraph(), c.profile)

        self._section(doc, "Competencias Técnicas")
        for category, items in c.skills:
            self._labelled(doc, category, items)

        self._section(doc, "Experiencia Profesional")
        for exp in c.experience:
            self._run(doc.add_paragraph(), f"{exp.company}  —  {exp.role}", bold=True, size=11, color=PRIMARY)
            self._run(doc.add_paragraph(), exp.period, size=9, color=MUTED, italic=True)
            for bullet in exp.bullets:
                self._bullet(doc, bullet)

        self._section(doc, "Educación")
        for degree, inst, year in c.education:
            p = doc.add_paragraph()
            self._run(p, degree, bold=True, size=10.5, color=PRIMARY)
            self._run(p, f"  —  {inst}  ({year})")

        self._section(doc, "Certificaciones & Formación")
        self._subsection(doc, "Certificaciones Oficiales")
        for line in c.official_certifications:
            self._bullet(doc, line)
        self._subsection(doc, "Rutas de Aprendizaje")
        for line in c.learning_paths:
            self._bullet(doc, line)
        self._subsection(doc, "Cursos de Formación")
        for category, lines in c.courses_by_category:
            self._run(doc.add_paragraph(), category, bold=True, color=PRIMARY)
            for line in lines:
                self._bullet(doc, line)

        self._section(doc, "Portafolio")
        p = doc.add_paragraph()
        self._run(p, "Sitio web: ", bold=True, color=PRIMARY)
        self._run(p, c.portfolio, color=ACCENT)

        self._section(doc, "Idiomas")
        for lang, level in c.languages:
            self._labelled(doc, lang, level)

    @staticmethod
    def _run(paragraph, text: str, *, bold=False, size=10.0, color=TEXT, italic=False) -> None:
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        run.font.size = Pt(size)
        run.font.color.rgb = color

    def _centered(self, doc: DocxDocument, text: str, **style) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(p, text, **style)

    def _labelled(self, doc: DocxDocument, label: str, value: str) -> None:
        p = doc.add_paragraph()
        self._run(p, f"{label}: ", bold=True, color=PRIMARY)
        self._run(p, value)

    def _bullet(self, doc: DocxDocument, text: str) -> None:
        self._run(doc.add_paragraph(style="List Bullet"), text)

    def _subsection(self, doc: DocxDocument, text: str) -> None:
        p = doc.add_paragraph()
        self._run(p, text, bold=True, size=11, color=ACCENT)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)

    def _section(self, doc: DocxDocument, text: str) -> None:
        p = doc.add_paragraph()
        self._run(p, text.upper(), bold=True, size=12, color=PRIMARY)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "12"), ("w:space", "1"), ("w:color", ACCENT_HEX)):
            bottom.set(qn(key), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)
